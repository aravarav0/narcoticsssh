"""Generate the official SIH 6-slide idea deck for SIH26231.
Follows the SIH idea format: Title, Proposed Solution, Technical Approach,
Feasibility & Viability, Impact & Benefits, Research & References.
Points/diagrams, no paragraphs. Export to PDF before uploading to the portal.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = r"c:\Users\shama\OneDrive\Documents\smartindiahackathon\SIH26231-idea-6slide.pptx"

NAVY = RGBColor(0x0D, 0x24, 0x33)
FOREST = RGBColor(0x0D, 0x5C, 0x47)
GOLD = RGBColor(0xE1, 0xB6, 0x53)
PAPER = RGBColor(0xF3, 0xF6, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTE = RGBColor(0xBF, 0xD3, 0xCC)
INK = RGBColor(0x14, 0x24, 0x20)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, l, t, w, h):
    return slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame


def para(tf, text, size, color, bold=False, first=False, align=PP_ALIGN.LEFT,
         space_after=8, bullet=False, level=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.level = level
    r = p.add_run()
    r.text = ("•  " + text) if bullet else text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return p


def footer(slide, idx):
    tf = box(slide, 0.6, 6.95, 12.1, 0.4)
    p = para(tf, "SIH26231  ·  Team [TEAM NAME]  ·  Digital Companion for Field Drug Testing",
             11, MUTE, first=True)
    n = box(slide, 12.2, 6.95, 0.9, 0.4)
    para(n, f"{idx} / 6", 11, MUTE, first=True, align=PP_ALIGN.RIGHT)


def accent_bar(slide):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.42), Inches(2.2), Inches(0.08))
    s.fill.solid(); s.fill.fore_color.rgb = GOLD
    s.line.fill.background()


def header(slide, kicker, title):
    tf = box(slide, 0.6, 0.45, 12.1, 0.5)
    para(tf, kicker, 14, GOLD, bold=True, first=True)
    tf2 = box(slide, 0.6, 0.8, 12.1, 0.7)
    para(tf2, title, 32, WHITE, bold=True, first=True)
    accent_bar(slide)


def content_slide(idx, kicker, title, columns):
    """columns: list of (heading, [bullets]) rendered side by side."""
    s = prs.slides.add_slide(BLANK)
    bg(s, NAVY)
    header(s, kicker, title)
    n = len(columns)
    gap = 0.4
    total_w = 12.1
    col_w = (total_w - gap * (n - 1)) / n
    x = 0.6
    for heading, bullets in columns:
        tf = box(s, x, 1.75, col_w, 5.0)
        para(tf, heading, 18, GOLD, bold=True, first=True, space_after=10)
        for b in bullets:
            if b.startswith(">"):
                para(tf, b[1:].strip(), 13, MUTE, bold=False, space_after=6, level=1)
            else:
                para(tf, b, 15, WHITE, bullet=True, space_after=8)
        x += col_w + gap
    footer(s, idx)
    return s


# ---------- Slide 1: Title ----------
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
from pptx.enum.shapes import MSO_SHAPE
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.4), SW, Inches(2.6))
band.fill.solid(); band.fill.fore_color.rgb = FOREST
band.line.fill.background()
tf = box(s, 0.7, 0.7, 12, 0.6)
para(tf, "SMART INDIA HACKATHON 2026", 16, GOLD, bold=True, first=True)
tf = box(s, 0.7, 1.25, 12, 0.6)
para(tf, "Problem Statement ID  ·  SIH26231", 18, WHITE, first=True)
tf = box(s, 0.7, 2.7, 12, 1.4)
para(tf, "Digital Companion for Field Drug Testing", 40, WHITE, bold=True, first=True)
para(tf, "A phone witness for colour-change kits — no new hardware", 20, PAPER)
tf = box(s, 0.7, 5.3, 12, 1.7)
para(tf, "Theme:  MedTech / HealthTech  (with tamper-evident evidence)", 16, WHITE, first=True, space_after=6)
para(tf, "PS Category:  Software", 16, WHITE, space_after=6)
para(tf, "Organisation:  Narcotics Control Bureau · Ministry of Home Affairs", 16, WHITE, space_after=6)
para(tf, "Team:  [TEAM NAME]      Team ID:  [TEAM ID]      Shiv Nadar University", 16, GOLD, bold=True)

# ---------- Slide 2: Proposed Solution ----------
content_slide(
    2, "PROPOSED SOLUTION", "Photograph the kit. Correct the light. Seal the proof.",
    [
        ("What it does", [
            "Photo of the colour-change kit WITH a reference card in the same frame.",
            "App corrects the room lighting using the card.",
            "Names the colour (+ hex) and calls it Positive / Negative / Inconclusive.",
            "Seals the photo: time, GPS, officer ID, SHA-256 + digital signature.",
        ]),
        ("Why it is unique", [
            "Uses the kit that already exists — zero new hardware.",
            "Refuses to guess: bad photo, glare, or no card -> Inconclusive.",
            "Tamper-evident record — change one pixel, the seal breaks.",
            "Explainable colour maths, not a black-box neural net.",
            "> Turns a subjective eyeball call into a repeatable measurement + verifiable record.",
        ]),
    ],
)

# ---------- Slide 3: Technical Approach ----------
content_slide(
    3, "TECHNICAL APPROACH", "Deterministic colour science + web cryptography",
    [
        ("Stack", [
            "Vite + React + TypeScript PWA — runs in any phone browser.",
            "Works offline; installable; distribute by URL / QR.",
            "WebCrypto ECDSA P-256 signature; SHA-256 image hash.",
            "IndexedDB key store; hash-chained local log.",
        ]),
        ("Pipeline (per photo)", [
            "Sample 6 card patches + kit box -> median RGB.",
            "Linearise (IEC 61966-2-1 sRGB).",
            "Fit 3x3 colour-correction matrix; fallback von Kries on gray.",
            "Convert to CIE Lab; ΔE76 to Positive / Negative / Muddy.",
            "Quality gates (glare, clipping, card-missing) -> Inconclusive.",
            "> Same photo -> same call, on any phone, under any light.",
        ]),
    ],
)

# ---------- Slide 4: Feasibility & Viability ----------
content_slide(
    4, "FEASIBILITY & VIABILITY", "Buildable today, cheap to scale",
    [
        ("Feasible now", [
            "Existing kits + existing phones + a printed card (~near zero cost).",
            "Working prototype built and demoed in under 24 hours.",
            "No procurement, no devices to ship.",
        ]),
        ("Risks & mitigation", [
            "Printer / paper colour drift -> in-app calibration saves YOUR card.",
            "Extreme lighting / glare -> gray-card white balance + reject to Inconclusive.",
            "Camera auto-white-balance -> card is in-frame every shot.",
            "Legal weight -> framed as presumptive only, lab confirmation required.",
            "> Scales via per-kit reference profiles and a shareable web link.",
        ]),
    ],
)

# ---------- Slide 5: Impact & Benefits ----------
content_slide(
    5, "IMPACT & BENEFITS", "Fair calls, real evidence, almost free",
    [
        ("For officers / NCB", [
            "Consistent field decisions — less officer-to-officer disagreement.",
            "Faster, clearer call at the point of check.",
            "Searchable, verifiable log of every test.",
        ]),
        ("Wider impact", [
            "Evidentiary: hash + signature + GPS + time = documentary record.",
            "Social: transparency and auditability reduce disputed calls.",
            "Economic: near-zero marginal cost; no hardware budget.",
            "Ethical: honest Inconclusive protects the accused and the officer.",
        ]),
    ],
)

# ---------- Slide 6: Research & References ----------
content_slide(
    6, "RESEARCH & REFERENCES", "Standards and sources",
    [
        ("Standards / methods", [
            "IEC 61966-2-1 — sRGB colour space & linearisation.",
            "CIE 1976 L*a*b* and ΔE76 colour difference.",
            "von Kries chromatic adaptation (white balance).",
            "SHA-256 (FIPS 180-4); ECDSA P-256 (FIPS 186) via WebCrypto.",
        ]),
        ("Domain / problem", [
            "NCB / MHA Problem Statement SIH26231 (sih.gov.in).",
            "Colorimetric presumptive drug tests (e.g. Marquis, Scott) — presumptive, need lab confirmation.",
            "Presumptive field tests require confirmatory laboratory analysis.",
            "> Prototype + colour maths write-up: our project repo (COLOUR-MATH.md).",
        ]),
    ],
)

prs.save(OUT)
print("wrote", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
