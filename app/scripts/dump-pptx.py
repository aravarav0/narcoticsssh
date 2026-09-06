import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1] if len(sys.argv) > 1 else r"c:\Users\shama\OneDrive\Documents\smartindiahackathon\SIH26231-SSH-pitch.pptx"
prs = Presentation(path)
print(f"slide size: {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in")
print(f"slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides, 1):
    print(f"\n===== SLIDE {i} =====")
    for shape in slide.shapes:
        kind = shape.shape_type
        if shape.has_text_frame:
            txt = "\n".join(p.text for p in shape.text_frame.paragraphs)
            if txt.strip():
                print(f"[text] {txt}")
        elif shape.has_table:
            print("[table]")
            for row in shape.table.rows:
                print("  | " + " | ".join(c.text for c in row.cells))
        elif kind == 13:
            print(f"[picture] {shape.name}")
        else:
            print(f"[shape:{kind}] {shape.name}")
